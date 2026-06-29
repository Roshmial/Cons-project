import app
import json
from pptx import Presentation

with app.db_connect() as conn:
    rows = conn.execute("SELECT * FROM messages WHERE thread_id = ? AND id <= ? ORDER BY id ASC", (242, 997)).fetchall()
    thread = conn.execute("SELECT * FROM threads WHERE id = ?", (242,)).fetchone()
    user = conn.execute("SELECT * FROM users WHERE id = ?", (thread['user_id'],)).fetchone()
    payload = app.process_chat_task_with_focused_context(rows, app.user_to_dict(user, conn), 'BI', {})

pptx_buf = app.build_message_export_pptx(payload['message_row'], 'BI')
prs = Presentation(pptx_buf)
summary = []
for idx, slide in enumerate(prs.slides, start=1):
    texts = []
    for shape in slide.shapes:
        if hasattr(shape, 'text') and shape.text:
            text = ' | '.join(part.strip() for part in shape.text.splitlines() if part.strip())
            if text:
                texts.append(text)
    summary.append({'slide': idx, 'text': ' || '.join(texts[:3])[:260]})
print(json.dumps({'slides': len(prs.slides), 'sample': summary[:10]}, ensure_ascii=False))

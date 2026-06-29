import app
import json
from pptx import Presentation

MESSAGE_ID = 998
with app.db_connect() as conn:
    row = conn.execute("SELECT * FROM messages WHERE id = ?", (MESSAGE_ID,)).fetchone()

pptx_buf = app.build_message_export_pptx(row, 'BI')
prs = Presentation(pptx_buf)
summary = []
for idx, slide in enumerate(prs.slides, start=1):
    texts = []
    for shape in slide.shapes:
        if hasattr(shape, 'text') and shape.text:
            text = ' | '.join(part.strip() for part in shape.text.splitlines() if part.strip())
            if text:
                texts.append(text)
    summary.append({'slide': idx, 'text': ' || '.join(texts[:3])[:260]})
print(json.dumps({'slides': len(prs.slides), 'sample': summary[:12]}, ensure_ascii=False))

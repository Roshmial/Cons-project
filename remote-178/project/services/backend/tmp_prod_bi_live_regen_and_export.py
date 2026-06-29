import app
import json
from pptx import Presentation

with app.db_connect() as conn:
    rows = conn.execute("SELECT * FROM messages WHERE thread_id = ? AND id <= ? ORDER BY id ASC", (242, 997)).fetchall()
    thread = conn.execute("SELECT * FROM threads WHERE id = ?", (242,)).fetchone()
    user = conn.execute("SELECT * FROM users WHERE id = ?", (thread['user_id'],)).fetchone()
    profile = app.user_to_dict(user, conn)

reply_text, meta = app.call_hermes_api(rows, profile, 'BI', {})
fake_row = {
    'id': -1,
    'thread_id': 242,
    'role': 'assistant',
    'content': reply_text,
    'created_at': app.now_iso(),
    'meta_json': json.dumps({'message_kind': 'chat_response', **meta}, ensure_ascii=False),
}
pptx_buf = app.build_message_export_pptx(fake_row, 'BI')
prs = Presentation(pptx_buf)
summary = []
for idx, slide in enumerate(prs.slides, start=1):
    texts = []
    for shape in slide.shapes:
        if hasattr(shape, 'text') and shape.text:
            text = ' | '.join(part.strip() for part in shape.text.splitlines() if part.strip())
            if text:
                texts.append(text)
    summary.append({'slide': idx, 'text': ' || '.join(texts[:3])[:260]})
print(json.dumps({'slides': len(prs.slides), 'focused': meta.get('focused_followup_context'), 'sample': summary[:10]}, ensure_ascii=False))

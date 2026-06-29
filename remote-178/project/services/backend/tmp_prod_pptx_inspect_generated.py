from pptx import Presentation
path = '/home/hermes/workspace/hermes-web-mvp-react-8793/services/backend/data/message_exports/53b1620505064a85ade9d085805cdee1-BI-2026-06-26_13-10-00.pptx'
prs = Presentation(path)
print('slides', len(prs.slides))
for idx, slide in enumerate(list(prs.slides)[:10], start=1):
    texts = []
    for shape in slide.shapes:
        if hasattr(shape, 'text') and shape.text:
            texts.append(shape.text.replace('\n', ' | '))
    print(f'slide_{idx}', ' || '.join(texts[:5]))

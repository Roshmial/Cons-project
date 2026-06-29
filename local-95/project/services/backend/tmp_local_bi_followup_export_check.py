import app
from pptx import Presentation

with app.db_connect() as conn:
    rows = conn.execute("SELECT * FROM messages WHERE thread_id = ? AND id <= ? ORDER BY id ASC", (242, 997)).fetchall()
    user = conn.execute("SELECT * FROM users WHERE id = ?", (3,)).fetchone()
    profile = app.user_to_dict(user, conn)

reply_text, meta = app.call_hermes_api(rows, profile, 'BI', {})
_, file_meta = app.build_generated_file_reply(
    assistant_message_id=999003,
    reply_text=reply_text,
    reply_meta=meta,
    thread_title='BI',
    export_format='pptx',
    created_at='2026-06-26T13:40:00+00:00',
)
attachment = (file_meta.get('attachments') or [])[0]
path = attachment['local_path']
prs = Presentation(path)
print('slides', len(prs.slides))
for idx, slide in enumerate(list(prs.slides)[:10], start=1):
    texts = []
    for shape in slide.shapes:
        if hasattr(shape, 'text') and shape.text:
            texts.append(shape.text.replace('\n', ' | '))
    print(f'slide_{idx}', ' || '.join(texts[:4]))

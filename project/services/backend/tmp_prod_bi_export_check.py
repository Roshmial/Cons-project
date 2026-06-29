import json
import app
from pptx import Presentation

with app.db_connect() as conn:
    row = conn.execute("SELECT * FROM app.messages WHERE thread_id = ? AND id = ?", (242, 998)).fetchone()
    thread = conn.execute("SELECT title FROM app.threads WHERE id = ?", (242,)).fetchone()
print('row_found', bool(row))
thread_title = thread['title'] if thread else 'BI'
_, meta = app.build_generated_file_reply(
    assistant_message_id=998,
    reply_text=row['content'],
    reply_meta=json.loads(row['meta_json'] or '{}'),
    thread_title=thread_title,
    export_format='pptx',
    created_at=row['created_at'],
)
path = meta['attachments'][0]['local_path']
print(path)
prs = Presentation(path)
print('slides', len(prs.slides))
for i, slide in enumerate(prs.slides, 1):
    print(f'--- SLIDE {i} ---')
    pics=tables=0
    texts=[]
    for shape in slide.shapes:
        if shape.shape_type == 13:
            pics += 1
        if getattr(shape, 'has_table', False):
            tables += 1
        text = getattr(shape, 'text', '')
        if text and text.strip():
            texts.append(text.strip())
    print('pictures', pics, 'tables', tables)
    for t in texts[:12]:
        print(t)
